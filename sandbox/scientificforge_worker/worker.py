"""Fixed-operation CPU scientific worker.

This module deliberately exposes scientific operations, not a Python execution
surface. It does not evaluate supplied source code, invoke a shell, use the
network, install packages, mutate source datasets, or execute notebooks.
"""

from __future__ import annotations

import csv
from hashlib import sha256
from importlib import metadata
import json
import math
import os
from pathlib import Path
import platform
import stat
from typing import Any


MAX_MATRIX_DIMENSION = 64
MAX_INLINE_VALUES = 100_000
MAX_SOURCE_ROWS = 100_000
MAX_BOOTSTRAP_DRAW_VALUES = 5_000_000


class ScientificWorkerError(RuntimeError):
    """A fixed scientific job could not be executed safely."""


def _numpy():
    try:
        import numpy as np
    except Exception as exc:
        raise ScientificWorkerError(
            "numpy_unavailable"
        ) from exc

    return np


def _version(distribution: str) -> str:
    try:
        return metadata.version(distribution)
    except metadata.PackageNotFoundError:
        return "unavailable"


def _engine_versions(
    *,
    used_openpyxl: bool = False,
) -> dict[str, str]:
    versions = {
        "python": platform.python_version(),
        "numpy": _version("numpy"),
    }

    if used_openpyxl:
        versions["openpyxl"] = _version(
            "openpyxl"
        )

    return versions


def _sha256_file(path: Path) -> str:
    digest = sha256()

    with path.open("rb") as handle:
        for chunk in iter(
            lambda: handle.read(1024 * 1024),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest()


def _verified_snapshot(
    path_text: str,
    expected_sha256: str,
) -> Path:
    path = Path(
        str(path_text or "")
    )

    if not path.is_absolute():
        raise ScientificWorkerError(
            "source_snapshot_must_be_absolute"
        )

    if path.is_symlink():
        raise ScientificWorkerError(
            "source_snapshot_symlink_not_allowed"
        )

    try:
        resolved = path.resolve(
            strict=True
        )
        info = resolved.stat()
    except OSError as exc:
        raise ScientificWorkerError(
            "source_snapshot_unavailable"
        ) from exc

    if not stat.S_ISREG(info.st_mode):
        raise ScientificWorkerError(
            "source_snapshot_regular_file_required"
        )

    if info.st_nlink != 1:
        raise ScientificWorkerError(
            "source_snapshot_hardlink_not_allowed"
        )

    observed = _sha256_file(
        resolved
    )

    if observed != expected_sha256:
        raise ScientificWorkerError(
            "source_snapshot_digest_mismatch"
        )

    return resolved


def _finite_float(value: Any) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ScientificWorkerError(
            "nonnumeric_value"
        ) from exc

    if not math.isfinite(result):
        raise ScientificWorkerError(
            "nonfinite_value"
        )

    return result


def _inline_values(
    values: list[Any],
) -> list[float]:
    if not values:
        raise ScientificWorkerError(
            "numeric_values_required"
        )

    if len(values) > MAX_INLINE_VALUES:
        raise ScientificWorkerError(
            "numeric_value_limit_exceeded"
        )

    return [
        _finite_float(value)
        for value in values
    ]


def _missing(value: Any) -> bool:
    text = (
        ""
        if value is None
        else str(value).strip()
    )

    return text.casefold() in {
        "",
        "na",
        "n/a",
        "nan",
        "none",
        "null",
        "-",
    }


def _csv_rows(
    path: Path,
) -> tuple[list[str], list[dict[str, Any]]]:
    rows: list[dict[str, Any]] = []

    with path.open(
        "r",
        encoding="utf-8-sig",
        errors="strict",
        newline="",
    ) as handle:
        reader = csv.DictReader(
            handle
        )
        columns = list(
            reader.fieldnames or []
        )

        if not columns:
            raise ScientificWorkerError(
                "table_header_required"
            )

        for index, row in enumerate(
            reader
        ):
            if index >= MAX_SOURCE_ROWS:
                raise ScientificWorkerError(
                    "source_row_limit_exceeded"
                )

            rows.append(
                dict(row)
            )

    return columns, rows


def _xlsx_rows(
    path: Path,
) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        import openpyxl
    except Exception as exc:
        raise ScientificWorkerError(
            "openpyxl_unavailable"
        ) from exc

    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook[
            workbook.sheetnames[0]
        ]

        iterator = worksheet.iter_rows(
            values_only=True
        )

        try:
            raw_header = next(
                iterator
            )
        except StopIteration as exc:
            raise ScientificWorkerError(
                "table_header_required"
            ) from exc

        columns = [
            str(value).strip()
            if value is not None
            else ""
            for value in raw_header
        ]

        if (
            not columns
            or any(
                not column
                for column in columns
            )
            or len(set(columns))
            != len(columns)
        ):
            raise ScientificWorkerError(
                "invalid_table_header"
            )

        rows: list[dict[str, Any]] = []

        for index, row in enumerate(
            iterator
        ):
            if index >= MAX_SOURCE_ROWS:
                raise ScientificWorkerError(
                    "source_row_limit_exceeded"
                )

            values = list(row)

            rows.append(
                {
                    column: (
                        values[position]
                        if position < len(values)
                        else None
                    )
                    for position, column in enumerate(
                        columns
                    )
                }
            )

        return columns, rows

    finally:
        workbook.close()


def _table_rows(
    path: Path,
) -> tuple[
    list[str],
    list[dict[str, Any]],
    bool,
]:
    suffix = path.suffix.lower()

    if suffix == ".csv":
        columns, rows = _csv_rows(
            path
        )
        return columns, rows, False

    if suffix == ".xlsx":
        columns, rows = _xlsx_rows(
            path
        )
        return columns, rows, True

    raise ScientificWorkerError(
        "unsupported_source_snapshot_type"
    )


def _selected_column(
    *,
    path: Path,
    column: str,
) -> tuple[list[float], bool]:
    columns, rows, used_openpyxl = (
        _table_rows(path)
    )

    if column not in columns:
        raise ScientificWorkerError(
            "selected_column_not_found"
        )

    values: list[float] = []

    for row in rows:
        raw = row.get(
            column
        )

        if _missing(raw):
            continue

        values.append(
            _finite_float(raw)
        )

    if not values:
        raise ScientificWorkerError(
            "selected_column_has_no_numeric_values"
        )

    return values, used_openpyxl


def _selected_matrix(
    *,
    path: Path,
    selected: list[str],
) -> tuple[Any, int, bool]:
    np = _numpy()

    if (
        len(selected) < 2
        or len(selected) > 16
        or len(set(selected)) != len(selected)
    ):
        raise ScientificWorkerError(
            "correlation_requires_2_to_16_unique_columns"
        )

    columns, rows, used_openpyxl = (
        _table_rows(path)
    )

    if any(
        column not in columns
        for column in selected
    ):
        raise ScientificWorkerError(
            "selected_column_not_found"
        )

    matrix: list[list[float]] = []

    for row in rows:
        raw_values = [
            row.get(column)
            for column in selected
        ]

        if any(
            _missing(value)
            for value in raw_values
        ):
            continue

        matrix.append(
            [
                _finite_float(value)
                for value in raw_values
            ]
        )

    if len(matrix) < 2:
        raise ScientificWorkerError(
            "correlation_requires_two_complete_rows"
        )

    return (
        np.asarray(
            matrix,
            dtype=float,
        ),
        len(matrix),
        used_openpyxl,
    )


def _summary(
    values: list[float],
) -> dict[str, Any]:
    np = _numpy()

    array = np.asarray(
        values,
        dtype=float,
    )

    if array.size == 0:
        raise ScientificWorkerError(
            "numeric_values_required"
        )

    return {
        "count": int(
            array.size
        ),
        "mean": float(
            np.mean(array)
        ),
        "median": float(
            np.median(array)
        ),
        "stddev_sample": (
            float(
                np.std(
                    array,
                    ddof=1,
                )
            )
            if array.size > 1
            else None
        ),
        "min": float(
            np.min(array)
        ),
        "max": float(
            np.max(array)
        ),
        "q25": float(
            np.quantile(
                array,
                0.25,
            )
        ),
        "q75": float(
            np.quantile(
                array,
                0.75,
            )
        ),
    }


def _descriptive_stats(
    job: dict[str, Any],
    source: Path | None,
) -> tuple[dict[str, Any], bool]:
    if source is not None:
        columns = list(
            job.get("columns") or []
        )

        if len(columns) != 1:
            raise ScientificWorkerError(
                "descriptive_stats_requires_one_source_column"
            )

        values, used_openpyxl = (
            _selected_column(
                path=source,
                column=str(
                    columns[0]
                ),
            )
        )

        return {
            "column": str(
                columns[0]
            ),
            "statistics": _summary(
                values
            ),
        }, used_openpyxl

    values = _inline_values(
        list(
            job.get("values")
            or []
        )
    )

    return {
        "statistics": _summary(
            values
        )
    }, False


def _correlation_matrix(
    job: dict[str, Any],
    source: Path | None,
) -> tuple[dict[str, Any], bool]:
    if source is None:
        raise ScientificWorkerError(
            "correlation_source_required"
        )

    np = _numpy()

    selected = [
        str(value)
        for value in (
            job.get("columns")
            or []
        )
    ]

    matrix, complete_rows, used_openpyxl = (
        _selected_matrix(
            path=source,
            selected=selected,
        )
    )

    correlation = np.corrcoef(
        matrix,
        rowvar=False,
    )

    normalized: list[list[float | None]] = []

    for row in correlation.tolist():
        normalized.append(
            [
                (
                    float(value)
                    if math.isfinite(
                        float(value)
                    )
                    else None
                )
                for value in row
            ]
        )

    return {
        "columns": selected,
        "complete_row_count": complete_rows,
        "matrix": normalized,
    }, used_openpyxl


def _bootstrap_mean_ci(
    job: dict[str, Any],
    source: Path | None,
) -> tuple[dict[str, Any], bool]:
    np = _numpy()

    seed = job.get(
        "seed"
    )

    if seed is None:
        raise ScientificWorkerError(
            "explicit_seed_required"
        )

    used_openpyxl = False

    if source is not None:
        columns = list(
            job.get("columns")
            or []
        )

        if len(columns) != 1:
            raise ScientificWorkerError(
                "bootstrap_requires_one_source_column"
            )

        values, used_openpyxl = (
            _selected_column(
                path=source,
                column=str(
                    columns[0]
                ),
            )
        )

    else:
        values = _inline_values(
            list(
                job.get("values")
                or []
            )
        )

    samples = int(
        job.get(
            "bootstrap_samples",
            2000,
        )
    )

    confidence = float(
        job.get(
            "confidence_level",
            0.95,
        )
    )

    if not 100 <= samples <= 10_000:
        raise ScientificWorkerError(
            "bootstrap_sample_limit"
        )

    if not 0.0 < confidence < 1.0:
        raise ScientificWorkerError(
            "invalid_confidence_level"
        )

    array = np.asarray(
        values,
        dtype=float,
    )

    if samples * int(array.size) > MAX_BOOTSTRAP_DRAW_VALUES:
        raise ScientificWorkerError(
            "bootstrap_work_limit_exceeded"
        )

    rng = np.random.default_rng(
        int(seed)
    )

    indexes = rng.integers(
        0,
        array.size,
        size=(
            samples,
            array.size,
        ),
    )

    means = np.mean(
        array[indexes],
        axis=1,
    )

    alpha = (
        1.0 - confidence
    ) / 2.0

    return {
        "count": int(
            array.size
        ),
        "observed_mean": float(
            np.mean(array)
        ),
        "bootstrap_samples": samples,
        "confidence_level": confidence,
        "mean_ci": [
            float(
                np.quantile(
                    means,
                    alpha,
                )
            ),
            float(
                np.quantile(
                    means,
                    1.0 - alpha,
                )
            ),
        ],
    }, used_openpyxl


def _matrix(
    value: Any,
    *,
    label: str,
) -> Any:
    np = _numpy()

    if (
        not isinstance(value, list)
        or not value
        or len(value) > MAX_MATRIX_DIMENSION
    ):
        raise ScientificWorkerError(
            f"{label}_dimension_invalid"
        )

    width: int | None = None
    rows: list[list[float]] = []

    for raw_row in value:
        if (
            not isinstance(raw_row, list)
            or not raw_row
            or len(raw_row) > MAX_MATRIX_DIMENSION
        ):
            raise ScientificWorkerError(
                f"{label}_dimension_invalid"
            )

        if width is None:
            width = len(
                raw_row
            )

        if len(raw_row) != width:
            raise ScientificWorkerError(
                f"{label}_not_rectangular"
            )

        rows.append(
            [
                _finite_float(item)
                for item in raw_row
            ]
        )

    return np.asarray(
        rows,
        dtype=float,
    )


def _matrix_multiply(
    job: dict[str, Any],
) -> tuple[dict[str, Any], bool]:
    left = _matrix(
        job.get(
            "matrix_a"
        ),
        label="matrix_a",
    )

    right = _matrix(
        job.get(
            "matrix_b"
        ),
        label="matrix_b",
    )

    if left.shape[1] != right.shape[0]:
        raise ScientificWorkerError(
            "matrix_dimension_mismatch"
        )

    product = left @ right

    return {
        "left_shape": [
            int(value)
            for value in left.shape
        ],
        "right_shape": [
            int(value)
            for value in right.shape
        ],
        "result_shape": [
            int(value)
            for value in product.shape
        ],
        "matrix": [
            [
                float(value)
                for value in row
            ]
            for row in product.tolist()
        ],
    }, False


def _monte_carlo_normal(
    job: dict[str, Any],
) -> tuple[dict[str, Any], bool]:
    np = _numpy()

    seed = job.get(
        "seed"
    )

    if seed is None:
        raise ScientificWorkerError(
            "explicit_seed_required"
        )

    samples = int(
        job.get(
            "monte_carlo_samples",
            10_000,
        )
    )

    if not 100 <= samples <= 100_000:
        raise ScientificWorkerError(
            "monte_carlo_sample_limit"
        )

    mean = _finite_float(
        job.get(
            "distribution_mean",
            0.0,
        )
    )

    stddev = _finite_float(
        job.get(
            "distribution_stddev",
            1.0,
        )
    )

    if stddev <= 0:
        raise ScientificWorkerError(
            "distribution_stddev_must_be_positive"
        )

    rng = np.random.default_rng(
        int(seed)
    )

    values = rng.normal(
        loc=mean,
        scale=stddev,
        size=samples,
    )

    return {
        "sample_count": samples,
        "configured_mean": mean,
        "configured_stddev": stddev,
        "observed_mean": float(
            np.mean(values)
        ),
        "observed_stddev_sample": float(
            np.std(
                values,
                ddof=1,
            )
        ),
        "q025": float(
            np.quantile(
                values,
                0.025,
            )
        ),
        "q50": float(
            np.quantile(
                values,
                0.5,
            )
        ),
        "q975": float(
            np.quantile(
                values,
                0.975,
            )
        ),
    }, False


def run_scientific_job(
    job: dict[str, Any],
) -> dict[str, Any]:
    """Execute one fixed ScientificForge operation."""

    if job.get("ir_version") == "scientific-ir-v0.2":
        if job.get("internal_backend_probe") is True and job.get("operation") == "backend_probe":
            # Only this fixed internal health contract accepts a backend family.
            # Every probe is launched by the parent in a fresh subprocess.
            from core.scientific_registry import BACKEND_MODULES
            import importlib
            from importlib.metadata import version
            family = str(job.get("backend_family") or "")
            if family not in {"highspy", "ortools"}:
                raise ScientificWorkerError("native_backend_family_not_allowed")
            module = BACKEND_MODULES[family]
            try:
                importlib.import_module(module)
                if family == "highspy":
                    from highspy import Highs
                    native_version = Highs().version()
                else:
                    from ortools.linear_solver import pywraplp
                    native_version = pywraplp.Solver.CreateSolver("GLOP").SolverVersion()
            except (ImportError, OSError) as exc:
                return {"status": "failed", "blocked_reason": "native_backend_import_failed:" + type(exc).__name__,
                        "network_access_used": False, "source_mutated": False,
                        "arbitrary_python_used": False, "shell_used": False,
                        "package_install_used": False}
            return {"status": "completed", "operation": "backend_probe",
                    "result": {"backend_family": family, "version": version(module), "native_version": native_version,
                               "process_id": os.getpid()},
                    "diagnostics": {"separate_process": True},
                    "network_access_used": False, "source_mutated": False,
                    "arbitrary_python_used": False, "shell_used": False,
                    "package_install_used": False}
        from sandbox.scientificforge_worker.typed_ops import (
            ScientificAdapterError,
            ScientificDependencyUnavailable,
            run_typed_operation,
        )
        try:
            return run_typed_operation(job)
        except ScientificDependencyUnavailable as exc:
            return {"status": "blocked", "blocked_reason": str(exc),
                    "network_access_used": False, "source_mutated": False,
                    "arbitrary_python_used": False, "shell_used": False,
                    "package_install_used": False}
        except (ScientificAdapterError, ArithmeticError) as exc:
            return {"status": "failed", "blocked_reason": str(exc)[:160],
                    "network_access_used": False, "source_mutated": False,
                    "arbitrary_python_used": False, "shell_used": False,
                    "package_install_used": False}

    operation = str(
        job.get(
            "operation"
        )
        or ""
    )

    source: Path | None = None

    source_snapshot = str(
        job.get(
            "source_snapshot"
        )
        or ""
    ).strip()

    expected_source_sha256 = str(
        job.get(
            "expected_source_sha256"
        )
        or ""
    ).strip()

    if source_snapshot:
        if not expected_source_sha256:
            raise ScientificWorkerError(
                "source_digest_required"
            )

        source = _verified_snapshot(
            source_snapshot,
            expected_source_sha256,
        )

    if operation == "descriptive_stats":
        result, used_openpyxl = (
            _descriptive_stats(
                job,
                source,
            )
        )

    elif operation == "correlation_matrix":
        result, used_openpyxl = (
            _correlation_matrix(
                job,
                source,
            )
        )

    elif operation == "bootstrap_mean_ci":
        result, used_openpyxl = (
            _bootstrap_mean_ci(
                job,
                source,
            )
        )

    elif operation == "matrix_multiply":
        if source is not None:
            raise ScientificWorkerError(
                "matrix_multiply_does_not_accept_source"
            )

        result, used_openpyxl = (
            _matrix_multiply(
                job
            )
        )

    elif operation == "monte_carlo_normal":
        if source is not None:
            raise ScientificWorkerError(
                "monte_carlo_does_not_accept_source"
            )

        result, used_openpyxl = (
            _monte_carlo_normal(
                job
            )
        )

    else:
        raise ScientificWorkerError(
            "unsupported_scientific_operation"
        )

    return {
        "status": "completed",
        "operation": operation,
        "result": result,
        "seed_used": (
            int(job["seed"])
            if job.get("seed")
            is not None
            else None
        ),
        "engine_versions": _engine_versions(
            used_openpyxl=used_openpyxl,
        ),
        "network_access_used": False,
        "source_mutated": False,
        "arbitrary_python_used": False,
        "shell_used": False,
        "package_install_used": False,
    }


__all__ = (
    "ScientificWorkerError",
    "run_scientific_job",
)
